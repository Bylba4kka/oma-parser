import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import defer

from app.db import Session
from app.models import Product, Status
from app.settings import settings

MAX_CELL = 32000
COLUMNS = [
    ("ID", 10),
    ("ID товара", 12),
    ("Артикул", 16),
    ("Название", 50),
    ("Ссылка на товар", 50),
    ("Бренд", 18),
    ("Категория", 40),
    ("Цена", 10),
    ("Старая цена", 12),
    ("Валюта", 8),
    ("Наличие", 14),
    ("Описание", 60),
    ("Характеристики (JSON)", 60),
    ("Фото", 40),
    ("Все фото", 40),
    ("Спарсено", 18),
]
LINK_COLUMNS = {"Ссылка на товар", "Фото"}
LINK_IDX = {i for i, (name, _) in enumerate(COLUMNS) if name in LINK_COLUMNS}


def _cell(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        return value[:MAX_CELL]
    return value


def _row(p: Product) -> list:
    params = json.dumps(p.params, ensure_ascii=False) if p.params else None
    return [
        p.id,
        p.external_id,
        p.sku,
        p.name,
        p.url,
        p.brand,
        p.category,
        float(p.price) if p.price is not None else None,
        float(p.old_price) if p.old_price is not None else None,
        p.currency,
        p.availability,
        p.description,
        params,
        p.image,
        "\n".join(p.images or []),
        p.parsed_at.replace(tzinfo=None) if p.parsed_at else None,
    ]


async def build_xlsx(only_done: bool = True, wide_params: int = 0) -> str:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Товары")

    stmt = (
        select(Product)
        .options(defer(Product.html), defer(Product.card))
        .order_by(Product.id)
        .execution_options(yield_per=2000)
    )
    if only_done:
        stmt = stmt.where(Product.status == Status.DONE)

    top_params: list[str] = []
    if wide_params:
        counter: Counter = Counter()
        async with Session() as s:
            async for params in await s.stream_scalars(
                select(Product.params).where(Product.status == Status.DONE).execution_options(yield_per=5000)
            ):
                counter.update((params or {}).keys())
        top_params = [k for k, _ in counter.most_common(wide_params)]

    for i, (_, width) in enumerate(COLUMNS + [("", 20)] * len(top_params)):
        ws.column_dimensions[get_column_letter(i + 1)].width = width
    ws.freeze_panes = "B2"
    header = [name for name, _ in COLUMNS] + top_params
    ws.append([_header_cell(ws, h) for h in header])

    async with Session() as s:
        async for p in await s.stream_scalars(stmt):
            row = _row(p)
            if top_params:
                params = p.params or {}
                row += [params.get(k) for k in top_params]
            ws.append([_link_cell(ws, v) if i in LINK_IDX and v else _cell(v) for i, v in enumerate(row)])

    path = export_dir() / f"oma_catalog_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    tmp = path.with_suffix(".part")
    wb.save(tmp)
    tmp.rename(path)
    return str(path)


def export_dir() -> Path:
    path = Path(settings.shared_dir) / "exports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _header_cell(ws, value: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.font = Font(bold=True)
    return cell


def _link_cell(ws, url: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=url)
    cell.hyperlink = url
    cell.style = "Hyperlink"
    return cell
